from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.exercise_generation import normalize_practice_set
from app.practice_export import build_practice_question_docx
from app.practice_question_format import OBJECTIVE_ANSWER_SLOT, ensure_objective_answer_slot


@pytest.mark.parametrize("suffix", ["（ ）。", "( )！", "（ ）。（      ）", "（ ） （ ）", "（\u3000）？"])
def test_punctuated_and_duplicated_answer_slots(suffix):
    result = ensure_objective_answer_slot("选择正确项" + suffix, "单选题")
    assert result.count("（") == 1
    assert ensure_objective_answer_slot(result, "单选题") == result
    assert ensure_objective_answer_slot("条件（x≠0），选择正确项" + suffix, "单选题").startswith("条件（x≠0），")


def test_word_export_repairs_punctuated_duplicate_slot():
    document = Document(BytesIO(build_practice_question_docx({"exercises": [{
        "number": 1, "question_type": "单选题", "stem": "选择正确项（ ）。（      ）",
        "options": [{"text": "甲"}, {"text": "乙"}],
    }]})))
    stems = [p.text for p in document.paragraphs if "选择正确项" in p.text]
    assert len(stems) == 1 and stems[0].count("（") == 1


def test_objective_answer_slot_is_exact_and_idempotent() -> None:
    assert ensure_objective_answer_slot("请选择正确结论。", "单选题") == f"请选择正确结论。{OBJECTIVE_ANSWER_SLOT}"
    assert ensure_objective_answer_slot("判断该说法。 (   )", "判断题") == f"判断该说法。{OBJECTIVE_ANSWER_SLOT}"
    assert ensure_objective_answer_slot(f"选择正确项。{OBJECTIVE_ANSWER_SLOT}", "选择题") == f"选择正确项。{OBJECTIVE_ANSWER_SLOT}"
    assert ensure_objective_answer_slot("说明原因。", "简答题") == "说明原因。"
    assert OBJECTIVE_ANSWER_SLOT.count(" ") == 6


def test_generated_choice_and_true_false_stems_are_saved_with_answer_slots() -> None:
    normalized = normalize_practice_set(
        {
            "exercises": [
                {"plan_item_id": "choice", "question_type": "单选题", "stem": "选择正确项。", "options": [{"label": "A", "text": "甲"}, {"label": "B", "text": "乙"}]},
                {"plan_item_id": "judge", "question_type": "判断题", "stem": "判断该说法。"},
            ]
        },
        requested_count=2,
        subject="材料科学",
        planned_plan_ids=["choice", "judge"],
    )

    assert [item["stem"] for item in normalized["exercises"]] == [
        f"选择正确项。{OBJECTIVE_ANSWER_SLOT}",
        f"判断该说法。{OBJECTIVE_ANSWER_SLOT}",
    ]


def test_word_export_repairs_older_objective_stems_without_duplicate_slots() -> None:
    data = {
        "source_analysis": {"subject": "材料科学"},
        "blueprint": {"training_goal": "客观题训练"},
        "exercises": [
            {"number": 1, "question_type": "多选题", "stem": "选择所有正确项。", "options": [{"text": "甲"}, {"text": "乙"}]},
            {"number": 2, "question_type": "判断题", "stem": "该命题成立。（ ）", "options": []},
        ],
    }

    document = Document(BytesIO(build_practice_question_docx(data)))
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]

    assert f"选择所有正确项。{OBJECTIVE_ANSWER_SLOT}" in paragraph_texts
    assert f"该命题成立。{OBJECTIVE_ANSWER_SLOT}" in paragraph_texts
    assert all(text.count("（") == 1 for text in paragraph_texts if OBJECTIVE_ANSWER_SLOT in text)
