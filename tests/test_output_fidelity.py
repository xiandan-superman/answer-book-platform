import json
from copy import deepcopy

from app.audit_model_repair import _repair_retry_prompt
from app.calculation_consistency import _close, formula_numeric_consistency_issues
from app.exercise_generation import _normalize_formulas, _normalize_generated_markup, _normalize_options, _normalize_tables
from app.expression_promotion import promote_inline_mathematical_expressions, promote_inline_reactions
from app.output_checkpoints import content_sha256
from app.practice_export import normalize_practice_markup


def test_step_inline_formula_and_adjacent_conditions_are_not_discarded():
    from docx import Document

    from app.docx_v4 import add_split_block

    doc = Document()
    add_split_block(doc, [
        {"type": "text", "text": "要得到 "},
        {"type": "formula_ref", "formula_id": "f", "inline": True},
        {"type": "text", "text": " 的聚合物"},
        {"type": "text", "text": "当 "},
        {"type": "formula_ref", "formula_id": "g", "inline": True},
        {"type": "text", "text": " 时"},
    ], {"f": {"latex": "Xn=200"}, "g": {"latex": "P=1"}})
    math = doc.element.xpath(".//m:oMath")
    assert len(math) == 2
    assert "Xn=200" in "".join(math[0].itertext())
    assert "P=1" in "".join(math[1].itertext())
    assert len(doc.paragraphs) == 2
    assert "的聚合物" in doc.paragraphs[0].text


def test_standalone_function_argument_promotes_without_equation_or_model_call():
    for text, expected in [("内耗 tan δ 随温度变化", r"\tan \delta"),
                           ("函数 sin θ 的值", r"\sin \theta")]:
        original = {"question_id": "q", "formulas": [],
                    "blocks": [{"label": "解析", "segments": [{"type": "text", "text": text}]}]}
        result = promote_inline_mathematical_expressions(original)
        assert [formula["latex"] for formula in result["formulas"]] == [expected]
        assert promote_inline_mathematical_expressions(deepcopy(result)) == result


def test_long_reaction_with_multiple_mathrm_groups_splits_without_losing_braces():
    from docx import Document

    from app.docx_v4 import _display_formula_lines, add_formula_paragraph

    left = r"\mathrm{CH_2-CH(OH)-CH_2-CH(OH)-CH_2-CH(OH)-CH_2-CH(OH)} + n\,\mathrm{HCHO}"
    right = r"\mathrm{CH_2-CH(O-CH_2-O)CH_2-CH} + n\,\mathrm{H_2O}"
    arrow = r"\xrightarrow{\mathrm{H}^+}"
    assert _display_formula_lines(left + arrow + right) == [left, "{}" + arrow + right]
    doc = Document()
    add_formula_paragraph(doc, left + arrow + right)
    assert doc.paragraphs[0].paragraph_format.keep_with_next
    assert not doc.paragraphs[1].paragraph_format.keep_with_next


def test_paired_emphasis_renders_as_word_style_without_deleting_unmatched_text():
    from docx import Document

    from app.docx_v4 import add_text_paragraph, append_domain_text_runs

    doc = Document()
    p = add_text_paragraph(doc, "前**重要条件**后**未闭合")
    assert p.text == "前重要条件后**未闭合"
    assert any(run.text == "重要条件" and run.bold for run in p.runs)
    p = doc.add_paragraph()
    append_domain_text_runs(p, "**有向量**和说明")
    assert p.text == "有向量和说明"


def test_summary_numeric_units_and_concentration_brackets_are_preserved():
    from app.docx_v4 import _answer_summary_formula_candidates

    source = "速率 Ri=2.09e-08 mol·L^-1·s^-1；浓度 [M·]=1.71e-08 mol·L^-1"
    candidates = _answer_summary_formula_candidates(source)
    rendered = " ".join(latex for _, _, latex in candidates)
    assert r"\times 10^{-8}" in rendered
    assert r"\mathrm{s}^{-1}" in rendered
    assert "[M·]" in rendered
    for start, end, _ in candidates:
        assert source[start:end]


def test_answer_summary_cleans_visible_percent_escape_without_mutating_delimited_formula(monkeypatch):
    from docx import Document
    from docx.oxml import OxmlElement
    import app.docx_v4 as docx_v4

    rendered = []

    def render(latex, **_kwargs):
        rendered.append(latex)
        return OxmlElement("m:oMath")

    monkeypatch.setattr(docx_v4, "render_expression_omml", render)
    doc = Document()
    paragraph = docx_v4.add_answer_summary_paragraph(
        doc,
        r"可见比例为63.2\%，公式为 $x=63.2\%$。",
    )

    assert r"\%" not in paragraph.text
    assert "63.2%" in paragraph.text
    assert rendered == [r"x=63.2\%"]


def test_word_formula_gate_counts_used_summary_but_not_stale_program_projection():
    from app.docx_v4 import expected_fragment_formula_count

    f = {"section": "计算题", "answer_summary": "value", "formulas": [
        {"formula_id": "model_formula", "latex": "x=1"},
        {"formula_id": "q_answer_summary_math_01", "latex": "y=2",
         "source_note": "程序在结构校验前从答案摘要中提升的数学表达式。"}],
         "blocks": [], "answer_summary_segments": [{"type": "formula_ref", "formula_id": "q_answer_summary_math_01"}]}
    assert expected_fragment_formula_count(f) == 1
    f["section"] = "简答题"
    assert expected_fragment_formula_count(f) == 2
    f["answer_summary_segments"] = []
    assert expected_fragment_formula_count(f) == 1
    f["blocks"] = [{"segments": [{"type": "formula_ref", "formula_id": "q_answer_summary_math_01"}]}]
    assert expected_fragment_formula_count(f) == 2


def test_repair_retains_original_evidence_and_images():
    messages = [
        {"role": "system", "content": "contract"},
        {"role": "user", "content": [
            {"type": "text", "text": "original question and evidence"},
            {"type": "image_url", "image_url": {"url": "data:image/png;base64,fixture"}},
        ]},
    ]
    before = deepcopy(messages)
    result = _repair_retry_prompt(messages, {"question_id": "q", "answer": "latest"}, ["failure"])
    assert result[:2] == messages
    result[1]["content"][0]["text"] = "mutated"
    assert messages == before
    assert "latest" in json.dumps(result)


def test_tiny_physical_quantities_do_not_use_dimensionless_absolute_tolerance():
    assert not _close(2.566738416e-5, 2.5667e-7)
    assert _close(2.566738416e-5, 2.57e-5)
    formulas = [
        {"role": "substitution", "latex": r"R_p = 1.76 \times 10^2\text{ L}\cdot\text{mol}^{-1}\cdot\text{s}^{-1} \times 8.53\text{ mol}\cdot\text{L}^{-1} \times 1.7097 \times 10^{-8}\text{ mol}\cdot\text{L}^{-1}"},
        {"role": "result", "latex": r"R_p = 2.5667 \times 10^{-7}\text{ mol}\cdot\text{L}^{-1}\cdot\text{s}^{-1}"},
    ]
    assert formula_numeric_consistency_issues(formulas)


def test_long_content_survives_generation_and_export():
    text = "甲" * 13000 + "必要条件"
    generated, unsafe = _normalize_generated_markup(text)
    assert generated == text
    assert not unsafe
    assert normalize_practice_markup(generated) == text


def test_visible_options_formulas_and_table_cells_are_not_truncated():
    long_text = "甲" * 4000 + "必要条件"
    assert _normalize_options([{"text": long_text}])[0]["text"] == long_text
    assert _normalize_formulas([{"latex": long_text, "caption": long_text}])[0]["latex"] == long_text
    rows = [[long_text] * 12 for _ in range(35)]
    normalized = _normalize_tables([{"headers": [long_text] * 12, "rows": rows}])[0]
    assert normalized["rows"] == rows
    assert normalized["headers"] == [long_text] * 12


def test_retry_diagnostics_identify_processed_candidate_not_raw_draft():
    draft = {"answer": "model output"}
    processed = {"answer": "processed output", "blocks": []}
    messages = _repair_retry_prompt([], draft, ["failed"], checked_candidate=processed)
    retry = json.loads(messages[-1]["content"])
    assert retry["previous_candidate"] == draft
    assert retry["checked_candidate"] == processed
    assert retry["validation_tool_result"]["meta"]["candidate_sha256"] == content_sha256(processed)


def test_identical_failed_repair_stops_and_retains_candidate(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from app import audit_model_repair as repair

    path = tmp_path / "answer_fragments.json"
    original = {"fragments": [{"question_id": "q", "answer": "accepted"}]}
    path.write_text(json.dumps(original))
    candidate = {"question_id": "q", "answer": "latest", "blocks": [], "formulas": []}
    calls = []

    def generate(*args, **kwargs):
        calls.append(args)
        return {"question_id": "q", "answer": "raw latest"}

    monkeypatch.setattr(repair, "_repair_prompt", lambda **kwargs: [{"role": "system", "content": "contract"}])
    monkeypatch.setattr(repair, "provider_model_supports_vision", lambda *args: False)
    monkeypatch.setattr(repair, "fragment_from_analysis_draft", lambda *args: deepcopy(candidate))
    monkeypatch.setattr(repair, "validate_v4_answer_fragment", lambda fragment: ["fixture_failure"])
    monkeypatch.setattr(repair, "_repair_regressions", lambda *args: [])
    report = repair.repair_fragments_with_model_for_audit(
        path, {"items": [{"question_id": "q", "stem": "question"}]}, [],
        selection_data=None, provider=SimpleNamespace(name="test", max_tokens=1000), model="test",
        audit_stage="answer_generation", audit_report={"issues": [{"question_id": "q", "code": "fixture_failure"}]},
        client=SimpleNamespace(chat_json_object=generate),
    )
    assert len(calls) == 2
    assert not report["changed"]
    assert json.loads(path.read_text()) == original
    history = report["validation_tool_results"]["q"]
    assert history[-1]["error"]["code"] == "ANSWER_REPAIR_NO_PROGRESS"
    assert not history[-1]["error"]["retryable"]
    assert list(tmp_path.glob("output_checkpoints/*/latest.json"))


def test_complete_tex_is_not_split_into_valid_looking_substrings():
    expressions = [r"\bar{X}_n = 2\nu", r"\frac{\pi}{c}", r"R_p = 1.76 \times 10^2", r"x = \frac{a}{b}",
                   r"\mathrm{S_2O_8^{2-} + SO_3^{2-} \rightarrow SO_4^{2-} + SO_4^{\bullet -}}"]
    for latex in expressions:
        fragment = {"question_id": "q", "formulas": [], "blocks": [
            {"label": "解析", "segments": [{"type": "text", "text": "关系为" + latex + "，结束。"}]}]}
        result = promote_inline_mathematical_expressions(promote_inline_reactions(fragment))
        assert len(result["formulas"]) == 1
        assert result["formulas"][0]["latex"] == latex
        assert not any("\\" in s.get("text", "") for s in result["blocks"][0]["segments"])
        assert promote_inline_mathematical_expressions(deepcopy(result)) == result
