from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from app.analysis_profiles import (
    EVIDENCE_BACKED_ANALYSIS,
    QUESTION_ONLY_ANALYSIS,
    TEXTBOOK_EVIDENCE_ONLY_ANALYSIS,
    analysis_generates_answers,
    analysis_uses_textbook_evidence,
    is_textbook_evidence_only,
    normalize_analysis_profile,
    sanitize_question_only_fragments,
)
from app.answer_coverage_audit import audit_answer_coverage
from app.audit_model_repair import _repair_prompt as build_audit_repair_prompt
from app.docx_model_repair import _repair_prompt as build_docx_repair_prompt
from app.docx_v4 import build_docx_from_fragments
from app.prompts import build_answer_draft_prompt
from app.review_export import build_question_review, write_question_review_csv
from app.task_read_model import build_exam_run


def _user_payload(messages: list[dict]) -> dict:
    content = messages[-1]["content"]
    if isinstance(content, list):
        content = content[0]["text"]
    return json.loads(content)


def test_analysis_profile_defaults_and_textbook_policy() -> None:
    assert normalize_analysis_profile(None) == EVIDENCE_BACKED_ANALYSIS
    assert normalize_analysis_profile(QUESTION_ONLY_ANALYSIS) == QUESTION_ONLY_ANALYSIS
    assert analysis_uses_textbook_evidence(EVIDENCE_BACKED_ANALYSIS)
    assert not analysis_uses_textbook_evidence(QUESTION_ONLY_ANALYSIS)
    assert analysis_uses_textbook_evidence(TEXTBOOK_EVIDENCE_ONLY_ANALYSIS)
    assert is_textbook_evidence_only(TEXTBOOK_EVIDENCE_ONLY_ANALYSIS)
    assert not analysis_generates_answers(TEXTBOOK_EVIDENCE_ONLY_ANALYSIS)
    coverage = audit_answer_coverage(
        {"items": [{"question_id": "q1", "section": "", "number": "1"}]},
        {"fragments": [{"question_id": "q1", "section": "", "number": "1", "answer": "答案"}]},
        require_evidence=False,
    )
    assert coverage["warnings"] == []


def test_question_only_prompt_is_derived_without_textbook_payload() -> None:
    messages = build_answer_draft_prompt(
        {"question_id": "q1", "number": "1", "stem": "解释测试概念。", "question_type": "名词解释"},
        [{"textbook": "不应出现", "evidence_text": "不应进入提示"}],
        include_textbook_evidence=False,
    )
    payload = _user_payload(messages)
    serialized = json.dumps(payload, ensure_ascii=False)
    assert payload["analysis_profile"] == QUESTION_ONLY_ANALYSIS
    assert "textbook_content" not in payload
    assert "不应进入提示" not in serialized
    assert "This profile does not run textbook indexing" in serialized
    assert "不得要求用户补教材" in messages[0]["content"]


def test_answer_prompt_includes_extracted_question_tables() -> None:
    messages = build_answer_draft_prompt(
        {
            "question_id": "q-table",
            "number": "1",
            "stem": "根据下表回答问题。",
            "question_type": "简答题",
            "attachments": {
                "tables": [
                    {"rows": [["项目", "数值"], ["A", "12"]], "text": "项目 | 数值\nA | 12"}
                ]
            },
        },
        [],
        include_textbook_evidence=False,
    )
    payload = _user_payload(messages)
    assert payload["question"]["tables"][0]["rows"][1] == ["A", "12"]
    assert "项目 | 数值" in payload["question"]["tables"][0]["text"]


def test_question_only_repair_prompts_do_not_restore_evidence_context() -> None:
    question = {"question_id": "q1", "number": "1", "stem": "解释测试概念。", "question_type": "名词解释"}
    audit_messages = build_audit_repair_prompt(
        audit_stage="content_quality",
        question=question,
        evidence=[{"evidence_text": "不应进入回修"}],
        fragment={"question_id": "q1", "answer": "待修复"},
        issues=[{"question_id": "q1", "code": "answer_too_short"}],
        include_textbook_evidence=False,
    )
    audit_payload = _user_payload(audit_messages)
    assert audit_payload["analysis_profile"] == QUESTION_ONLY_ANALYSIS
    assert audit_payload["confirmed_evidence"] == []
    assert "不应进入回修" not in json.dumps(audit_payload, ensure_ascii=False)

    docx_messages = build_docx_repair_prompt(
        question,
        [{"evidence_text": "不应进入格式回修"}],
        {"question_id": "q1", "answer": "待修复", "blocks": []},
        [{"question_id": "q1", "message": "测试问题"}],
        ["测试问题"],
        include_textbook_evidence=False,
    )
    docx_payload = _user_payload(docx_messages)
    assert docx_payload["analysis_profile"] == QUESTION_ONLY_ANALYSIS
    assert docx_payload["confirmed_evidence"] == []
    assert "不应进入格式回修" not in json.dumps(docx_payload, ensure_ascii=False)


def test_question_only_document_title_and_public_task_label(tmp_path: Path) -> None:
    source = tmp_path / "fragments.json"
    output = tmp_path / "result.docx"
    source.write_text(
        json.dumps(
            {
                "document_title": "题目解析",
                "analysis_profile": QUESTION_ONLY_ANALYSIS,
                "fragments": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    build_docx_from_fragments(source, output)
    assert Document(output).paragraphs[0].text == "题目解析"

    task = build_exam_run(
        {
            "task_id": "task-1",
            "exam_path": "/tmp/example.docx",
            "provider": "demo",
            "model": "demo-model",
            "status": "created",
            "current_stage": "created",
            "error": "",
            "created_at": "2026-08-29 00:00:00",
            "updated_at": "2026-08-29 00:00:00",
            "analysis_profile": QUESTION_ONLY_ANALYSIS,
            "selected_textbooks": ["/tmp/stale-textbook.pdf"],
        }
    )
    assert task["analysis_profile"] == QUESTION_ONLY_ANALYSIS
    assert task["display_title"].startswith("题目解析")
    assert task["textbook_material_names"] == []


def test_question_only_sanitizes_stale_textbook_evidence_from_data_and_docx(tmp_path: Path) -> None:
    payload = {
        "schema_version": "answer_book.answer_fragments.v4",
        "analysis_profile": QUESTION_ONLY_ANALYSIS,
        "document_title": "真题答案解析",
        "recovery_events": [{"question_id": "q1", "strategy": "program_evidence_binding"}],
        "fragments": [
            {
                "schema_version": "answer_book.answer_fragment.v4",
                "question_id": "q1",
                "section": "一、简答题",
                "question_type": "简答题",
                "number": "1",
                "answer": "答案正文。",
                "answer_summary": "答案正文。",
                "evidence_ids": ["e1"],
                "warnings": ["程序自动绑定教材依据", "保留提示"],
                "_meta": {"evidence_binding": {"strategy": "program_top_evidence"}, "other": True},
                "formulas": [],
                "blocks": [
                    {"label": "教材依据", "segments": [{"type": "text", "text": "不应出现的教材内容"}]},
                    {"label": "解析", "segments": [{"type": "text", "text": "解析正文。"}]},
                ],
            }
        ],
    }
    sanitized = sanitize_question_only_fragments(payload)
    fragment = sanitized["fragments"][0]
    assert sanitized["document_title"] == "题目解析"
    assert fragment["evidence_ids"] == []
    assert [block["label"] for block in fragment["blocks"]] == ["解析"]
    assert fragment["warnings"] == ["保留提示"]
    assert fragment["_meta"] == {"other": True}
    assert sanitized["recovery_events"] == []

    # Defense in depth: the Word builder must also suppress a stale evidence
    # block even when called directly with an unsanitized checkpoint file.
    payload["fragments"][0]["blocks"].insert(
        0,
        {"label": "教材依据", "segments": [{"type": "text", "text": "仍不应进入 Word"}]},
    )
    payload["fragments"][0]["evidence_ids"] = ["e2"]
    source = tmp_path / "question-only.json"
    output = tmp_path / "question-only.docx"
    source.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    build_docx_from_fragments(source, output)
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "题目解析" in text
    assert "教材依据" not in text
    assert "仍不应进入 Word" not in text


def test_question_only_review_omits_stale_evidence_columns_and_candidates(tmp_path: Path) -> None:
    (tmp_path / "structured_exam.json").write_text(
        json.dumps({"items": [{"question_id": "q1", "number": "1", "stem": "测试题"}]}, ensure_ascii=False),
        encoding="utf-8",
    )
    (tmp_path / "answer_fragments.json").write_text(
        json.dumps(
            {
                "analysis_profile": QUESTION_ONLY_ANALYSIS,
                "fragments": [
                    {
                        "question_id": "q1",
                        "answer": "答案",
                        "evidence_ids": ["e1"],
                        "_meta": {"evidence_binding": {"strategy": "program_top_evidence"}},
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (tmp_path / "retrieval_candidates.csv").write_text(
        "question_id,evidence_id,textbook\nq1,e1,不应读取的教材\n",
        encoding="utf-8-sig",
    )
    review = build_question_review(tmp_path)
    row = review["review_rows"][0]
    assert review["uses_textbook_evidence"] is False
    assert review["auto_evidence_count"] == 0
    assert row["evidence_id_count"] == 0
    assert row["candidate_count"] == 0
    assert row["top_candidates"] == []

    output = write_question_review_csv(review, tmp_path / "question_review.csv")
    header = output.read_text(encoding="utf-8-sig").splitlines()[0]
    assert "evidence" not in header
    assert "candidate" not in header


def test_question_only_analysis_uses_shared_manual_structure_review_gate() -> None:
    import inspect

    from app import pipeline

    source = inspect.getsource(pipeline._run_pipeline_impl)
    review_call = source.index("wait_for_exam_structure_review(")
    question_only_textbook_skip = source.index('if not textbook_evidence_enabled:')

    assert review_call < question_only_textbook_skip
    assert "auto_confirm_exam_structure" not in source


def test_question_only_entry_and_request_contract_are_present() -> None:
    root = Path(__file__).resolve().parents[1]
    index = (root / "web" / "index.html").read_text(encoding="utf-8")
    app = (root / "web" / "app.js").read_text(encoding="utf-8")
    assert 'id="questionAnalysisUtilityTitle">题目解析' in index
    assert 'id="reasoningModelRoleCard"' in index
    assert "function startQuestionAnalysis()" in app
    assert 'currentExamAnalysisProfile = "question_only";\n  goToPage("env");' in app
    assert '$("reasoningModelRoleCard")?.classList.toggle("hidden", questionOnly);' in app
    assert "function examRequiredTextRoutes()" in app
    assert "analysis_profile: currentExamAnalysisProfile" in app
    assert "if (!questionOnly) await requirePreparedTextbookIndex();" in app
    assert 'let activeTaskAnalysisProfile = "evidence_backed";' in app
    assert 'if (page !== "task") stopTaskPolling();' in app
    assert 'currentPage !== "task" || activeTaskId !== taskId' in app
    assert 'currentExamAnalysisProfile = task.analysis_profile' not in app
    assert '["result-question-evidence", "教材引用", !questionOnly]' in app
